"""Generic, formula-driven three-statement model builder.

Input: a model spec (dict, see SPEC.md) with quarterly history, forecast assumptions,
competitor data, industry series and source links.
Output: an .xlsx workbook with tabs Quarterly Model, Annual Model, Key Metrics,
Competitors and Data Sources, plus a validation report.
"""
from __future__ import annotations

import datetime as dt
import re
from dataclasses import dataclass, field

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

# ------------------------------------------------------------------ styles
FONT = "Arial"
f_norm = Font(name=FONT, size=10)
f_bold = Font(name=FONT, size=10, bold=True)
f_in = Font(name=FONT, size=10, color="0000FF")
f_link = Font(name=FONT, size=10, color="008000")
f_title = Font(name=FONT, size=14, bold=True)
f_sec = Font(name=FONT, size=11, bold=True, color="FFFFFF")
f_hdr = Font(name=FONT, size=10, bold=True)
f_small = Font(name=FONT, size=9, italic=True, color="7F7F7F")
f_url = Font(name=FONT, size=9, color="0563C1", underline="single")
f_chk = Font(name=FONT, size=9, italic=True, color="375623")
fill_sec = PatternFill("solid", fgColor="1F3864")
fill_fc = PatternFill("solid", fgColor="EEF3FA")
fill_inp = PatternFill("solid", fgColor="FFF2CC")
fill_chk = PatternFill("solid", fgColor="E2EFDA")
fill_th = PatternFill("solid", fgColor="D9E1F2")
thin_top = Border(top=Side(style="thin"))
NUM = '#,##0.0;(#,##0.0);"-"'
NUM0 = '#,##0;(#,##0);"-"'
PCT = '0.0%;(0.0%);"-"'
PCT2 = '0.00%;(0.00%);"-"'
EPS = '$0.00;($0.00);"-"'
DAYS = '0.0;(0.0);"-"'
MULT = '0.0x;(0.0x);"-"'
FC0 = 3  # first period column on period sheets


class SpecError(ValueError):
    pass


# ------------------------------------------------------------------ helpers
def _q_next(fy: int, fq: int):
    return (fy + 1, 1) if fq == 4 else (fy, fq + 1)


def _add_months(d: dt.date, months: int) -> dt.date:
    m = d.month - 1 + months
    y = d.year + m // 12
    m = m % 12 + 1
    # last day of month
    nxt = dt.date(y + (m == 12), m % 12 + 1, 1)
    return nxt - dt.timedelta(days=1)


def _num(v):
    if v is None:
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _series(d: dict, key: str, n: int, required=False, default=0.0):
    v = d.get(key)
    if v is None:
        if required:
            raise SpecError(f"missing required series '{key}'")
        return [default] * n if default is not None else [None] * n
    if len(v) != n:
        raise SpecError(f"series '{key}' has {len(v)} values; expected {n} (one per historical quarter)")
    out = [_num(x) for x in v]
    if default is not None:
        out = [default if x is None else x for x in out]
    return out


def _urls(text):
    if not isinstance(text, str):
        return []
    return [u.rstrip(".,;)") for u in re.findall(r"https?://[^\s;|,]+", text)]


@dataclass
class Report:
    errors: list = field(default_factory=list)
    warnings: list = field(default_factory=list)
    checks: dict = field(default_factory=dict)

    def as_dict(self):
        return {"ok": not self.errors, "errors": self.errors, "warnings": self.warnings, "checks": self.checks}


# ------------------------------------------------------------------ builder
class ModelBuilder:
    def __init__(self, spec: dict):
        self.spec = spec
        self.rep = Report()
        self._parse()

    # ---------------- spec parsing
    def _parse(self):
        s = self.spec
        h = s.get("history") or {}
        per = h.get("periods") or []
        if len(per) < 4:
            raise SpecError("history.periods needs at least 4 quarters")
        self.hist = []
        for p in per:
            try:
                self.hist.append({"label": p["label"], "end": dt.date.fromisoformat(p["end"]), "fy": int(p["fy"]), "fq": int(p["fq"])})
            except Exception as e:  # noqa: BLE001
                raise SpecError(f"bad period entry {p}: {e}")
        for a, b in zip(self.hist, self.hist[1:]):
            if _q_next(a["fy"], a["fq"]) != (b["fy"], b["fq"]):
                raise SpecError(f"periods are not consecutive fiscal quarters: {a['label']} -> {b['label']}")
        self.NH = len(self.hist)
        end_fy = int((s.get("forecast") or {}).get("end_fy") or (self.hist[-1]["fy"] + 5))
        self.periods = list(self.hist)
        fy, fq, d = self.hist[-1]["fy"], self.hist[-1]["fq"], self.hist[-1]["end"]
        while True:
            fy, fq = _q_next(fy, fq)
            if fy > end_fy:
                break
            d = _add_months(d, 3)
            self.periods.append({"label": f"Q{fq} FY{fy}", "end": d, "fy": fy, "fq": fq})
        self.NP = len(self.periods)
        # fiscal years shown on the annual sheet: only complete years (all 4 quarters present)
        years = sorted({p["fy"] for p in self.periods})
        self.years = [y for y in years if sum(1 for p in self.periods if p["fy"] == y) == 4]
        if not self.years:
            raise SpecError("no complete fiscal year in history + forecast")

        n = self.NH
        isd = h.get("income_statement") or {}
        bsd = h.get("balance_sheet") or {}
        cfd = h.get("cash_flow") or {}
        self.is_ = {
            "revenue": _series(isd, "revenue", n, required=True),
            "cost_of_revenue": _series(isd, "cost_of_revenue", n, required=True),
            "da": _series(isd, "depreciation_amortization", n),
            "interest": _series(isd, "interest_and_other", n),
            "tax": _series(isd, "income_tax", n, required=True),
            "ni_rep": _series(isd, "net_income_reported", n, default=None),
            "shares": _series(isd, "diluted_shares", n, required=True),
            "dps": _series(isd, "dps", n),
            "eps_rep": _series(isd, "eps_reported", n, default=None),
            "adj": _series(isd, "adjustments", n),
        }
        self.da_in_opex = bool(isd.get("da_in_opex", True))
        self.gross_label = isd.get("gross_profit_label") or "Gross profit"
        self.cogs_label = isd.get("cost_of_revenue_label") or "Cost of revenue"
        self.opex = []
        for o in isd.get("opex") or []:
            drv = o.get("driver", "pct_revenue")
            if drv not in ("pct_revenue", "pct_gross_profit", "headcount"):
                raise SpecError(f"opex '{o.get('name')}' has unknown driver '{drv}'")
            self.opex.append({"name": o["name"], "values": _series(o, "values", n, required=True), "driver": drv})
        if not self.opex:
            raise SpecError("income_statement.opex needs at least one operating expense line")
        self.headcount = _series(h, "headcount", n, default=None) if h.get("headcount") else None
        if self.headcount is None or any(v is None for v in self.headcount):
            for o in self.opex:
                if o["driver"] == "headcount":
                    self.rep.warnings.append(f"opex '{o['name']}' uses the headcount driver but headcount history is incomplete; switched to % of gross profit")
                    o["driver"] = "pct_gross_profit"
            self.headcount = None

        segs = h.get("segments") or []
        self.segs = []
        for sg in segs:
            self.segs.append({
                "name": sg["name"],
                "revenue": _series(sg, "revenue", n, required=True),
                "gp": _series(sg, "gross_profit", n, default=None) if sg.get("gross_profit") else None,
                "vol": _series(sg, "volume_growth", n, default=None) if sg.get("volume_growth") else [None] * n,
                "vol_label": sg.get("volume_label") or "volume",
                "price_label": sg.get("price_label") or "price / mix",
            })
        if not self.segs:
            gp = [r - c for r, c in zip(self.is_["revenue"], self.is_["cost_of_revenue"])]
            self.segs = [{"name": "Total company", "revenue": self.is_["revenue"], "gp": gp, "vol": [None] * n,
                          "vol_label": "volume", "price_label": "price / mix"}]
        # segment GP: if any segment lacks GP, allocate company GP pro rata to revenue
        tot_gp = [r - c for r, c in zip(self.is_["revenue"], self.is_["cost_of_revenue"])]
        if any(sg["gp"] is None for sg in self.segs):
            self.rep.warnings.append("segment gross profit not provided for every segment; allocated company gross profit by revenue share")
            for i in range(n):
                tr = sum(sg["revenue"][i] for sg in self.segs) or 1.0
                for sg in self.segs:
                    if sg["gp"] is None:
                        sg["gp_alloc"] = True
                    sg.setdefault("_gp", [0.0] * n)[i] = tot_gp[i] * sg["revenue"][i] / tr
            for sg in self.segs:
                sg["gp"] = sg.pop("_gp")
        for i in range(n):
            d_rev = sum(sg["revenue"][i] for sg in self.segs) - self.is_["revenue"][i]
            d_gp = sum(sg["gp"][i] for sg in self.segs) - tot_gp[i]
            if abs(d_rev) > 0.5 or abs(d_gp) > 0.5:
                self.rep.errors.append(f"{self.hist[i]['label']}: segments do not sum to company totals (revenue diff {d_rev:,.1f}, gross profit diff {d_gp:,.1f})")

        bs_keys = ["cash", "receivables", "inventory", "other_current_assets", "ppe_intangibles", "goodwill", "other_noncurrent_assets",
                   "payables", "other_current_liabilities", "short_term_debt", "long_term_debt", "other_noncurrent_liabilities", "equity"]
        self.bs = {k: _series(bsd, k, n, required=k in ("cash", "equity")) for k in bs_keys}
        self.bs_ta_rep = _series(bsd, "total_assets_reported", n, default=None) if bsd.get("total_assets_reported") else None
        ob = h.get("opening_balance") or {}
        self.opening = {k: _num(ob.get(k)) for k in bs_keys}
        cf_keys = ["sbc", "other_noncash", "change_receivables", "change_inventory", "change_payables", "change_other_wc", "cfo_reported",
                   "capex", "acquisitions", "other_investing", "cfi_reported", "buybacks", "dividends", "debt_net", "other_financing",
                   "cff_reported", "fx"]
        self.cf = {k: _series(cfd, k, n, required=k in ("cfo_reported", "capex", "cfi_reported", "cff_reported")) for k in cf_keys}

        # validations
        for i in range(n):
            a = sum(self.bs[k][i] for k in bs_keys[:7])
            le = sum(self.bs[k][i] for k in bs_keys[7:])
            if abs(a - le) > 0.5:
                self.rep.errors.append(f"{self.hist[i]['label']}: balance sheet does not balance (assets {a:,.1f} vs liabilities + equity {le:,.1f}); add the missing line items to the 'other' buckets")
            if self.bs_ta_rep and self.bs_ta_rep[i] is not None and abs(a - self.bs_ta_rep[i]) > 0.5:
                self.rep.errors.append(f"{self.hist[i]['label']}: asset lines sum to {a:,.1f} but reported total assets are {self.bs_ta_rep[i]:,.1f}")
            if self.is_["ni_rep"][i] is not None:
                ebit = self.is_["revenue"][i] - self.is_["cost_of_revenue"][i] - sum(o["values"][i] for o in self.opex) - (0 if self.da_in_opex else self.is_["da"][i])
                ni = ebit + self.is_["interest"][i] - self.is_["tax"][i]
                if abs(ni - self.is_["ni_rep"][i]) > 0.5:
                    self.rep.errors.append(f"{self.hist[i]['label']}: income statement lines give net income {ni:,.1f} vs reported {self.is_['ni_rep'][i]:,.1f}")
        self.plugs = {"cfo": [], "cfi": [], "cff": []}
        for i in range(n):
            ni = self.is_["revenue"][i] - self.is_["cost_of_revenue"][i] - sum(o["values"][i] for o in self.opex) - (0 if self.da_in_opex else self.is_["da"][i]) + self.is_["interest"][i] - self.is_["tax"][i]
            cfo_parts = ni + self.is_["da"][i] + sum(self.cf[k][i] for k in ["sbc", "other_noncash", "change_receivables", "change_inventory", "change_payables", "change_other_wc"])
            self.plugs["cfo"].append(self.cf["cfo_reported"][i] - cfo_parts)
            self.plugs["cfi"].append(self.cf["cfi_reported"][i] - sum(self.cf[k][i] for k in ["capex", "acquisitions", "other_investing"]))
            self.plugs["cff"].append(self.cf["cff_reported"][i] - sum(self.cf[k][i] for k in ["buybacks", "dividends", "debt_net", "other_financing"]))
        for k, vals in self.plugs.items():
            big = [(self.hist[i]["label"], v) for i, v in enumerate(vals) if abs(v) > max(5.0, 0.05 * abs(self.cf[k + "_reported"][i]))]
            if big:
                self.rep.warnings.append(f"{k.upper()} needed a reconciling 'other' line above 5% in {len(big)} quarter(s), e.g. {big[0][0]}: {big[0][1]:,.1f}. Check the cash-flow line mapping.")

    # ---------------- assumption resolution
    def _assume(self, key, p, default):
        a = (self.spec.get("forecast") or {}).get("assumptions") or {}
        d = a.get(key)
        if isinstance(d, (int, float)):
            return float(d)
        if isinstance(d, dict):
            for k in (f"{p['fy']}Q{p['fq']}", str(p["fy"])):
                if k in d and d[k] is not None:
                    return float(d[k])
        return default(p) if callable(default) else default

    def _trail(self, fn, n=4):
        vals = []
        for i in range(self.NH - n, self.NH):
            try:
                v = fn(i)
            except ZeroDivisionError:
                v = None
            if v is not None:
                vals.append(v)
        return sum(vals) / len(vals) if vals else 0.0

    def _same_q_last_year(self, fn):
        base = {}
        for i in range(max(0, self.NH - 4), self.NH):
            try:
                base[self.hist[i]["fq"]] = fn(i)
            except ZeroDivisionError:
                base[self.hist[i]["fq"]] = 0.0
        return lambda p: round(base.get(p["fq"], 0.0), 4)

    # ---------------- workbook
    def build(self, path: str) -> dict:
        wb = Workbook()
        self.wb = wb
        self._quarterly(wb.active)
        self._annual(wb.create_sheet("Annual Model"))
        self._metrics(wb.create_sheet("Key Metrics"))
        self._competitors(wb.create_sheet("Competitors"))
        self._sources(wb.create_sheet("Data Sources"))
        for ws in wb.worksheets:
            ws.sheet_view.showGridLines = False
        wb.save(path)
        return self.rep.as_dict()

    # ----- quarterly sheet engine
    def _quarterly(self, ws):
        ws.title = "Quarterly Model"
        self.ws = ws
        self.rows = {}
        self.cur = 7
        s = self.spec
        NC = FC0 + self.NP - 1
        unit = s.get("units", "millions")
        cur = s.get("currency", "USD")
        ws["A1"] = f"{s.get('company', s.get('ticker'))} ({s.get('exchange', '')}{': ' if s.get('exchange') else ''}{s.get('ticker')}) - Quarterly Three-Statement Model"
        ws["A1"].font = f_title
        fye = self.hist[-1]["end"] if self.hist[-1]["fq"] == 4 else _add_months(self.hist[-1]["end"], 3 * (4 - self.hist[-1]["fq"]))
        ws["A2"] = (f"{cur} {unit} unless noted. Fiscal year ends {fye.strftime('%B %d').replace(' 0', ' ')}; each column shows the fiscal quarter and its period-end date. "
                    "Blue = hard-coded input or reported actual; black = formula; green = link to another tab; yellow cells = forecast assumptions to edit.")
        ws["A3"] = f"Actuals {self.hist[0]['label']} - {self.hist[-1]['label']} from company filings (see Data Sources tab). Forecast {self.periods[self.NH]['label'] if self.NP > self.NH else '-'} - {self.periods[-1]['label']}."
        for c in ("A2", "A3"):
            ws[c].font = Font(name=FONT, size=9, italic=True)
        for i, p in enumerate(self.periods):
            c = FC0 + i
            ws.cell(4, c, p["label"]).font = f_hdr
            d = ws.cell(5, c, p["end"]); d.number_format = "m/d/yyyy"; d.font = f_norm
            ws.cell(6, c, "Actual" if i < self.NH else "Forecast").font = Font(name=FONT, size=9, italic=True, color="7F7F7F" if i < self.NH else "C00000")
            for r in (4, 5, 6):
                ws.cell(r, c).alignment = Alignment(horizontal="right")
                if i >= self.NH:
                    ws.cell(r, c).fill = fill_fc
        ws.cell(4, 1, "Fiscal quarter").font = f_hdr
        ws.cell(5, 1, "Period end date").font = f_hdr
        ws.cell(6, 1, "Actual / Forecast").font = f_hdr
        # two passes so formulas can reference rows defined later
        self._pass = 1
        self._layout(ws, NC)
        self._pass = 2
        for row in ws.iter_rows(min_row=7, max_row=ws.max_row):
            for c in row:
                c.value = None; c.comment = None; c.font = f_norm; c.fill = PatternFill(); c.border = Border()
        self.cur = 7
        self._layout(ws, NC)
        ws.column_dimensions["A"].width = 60
        ws.column_dimensions["B"].width = 7
        for i in range(self.NP):
            ws.column_dimensions[L(FC0 + i)].width = 11
        ws.freeze_panes = "C7"

    def C(self, i):
        return L(FC0 + i)

    def R(self, key, i, off=0):
        return f"{self.C(i + off)}{self.rows.get(key, 1)}"

    def section(self, title, ncols):
        r = self.cur; self.cur += 1
        self.ws.cell(r, 1, title).font = f_sec
        for c in range(1, ncols + 1):
            self.ws.cell(r, c).fill = fill_sec

    def blank(self):
        self.cur += 1

    def row(self, key, label, unit, hist, fcst, fmt=NUM, bold=False, top=False, note=None, check=False):
        r = self.cur; self.cur += 1
        self.rows[key] = r
        ws = self.ws
        ws.cell(r, 1, label).font = f_chk if check else (f_bold if bold else f_norm)
        ws.cell(r, 2, unit).font = f_small
        if note and self._pass == 2:
            ws.cell(r, 1).comment = Comment(note, "Model")
        for i, p in enumerate(self.periods):
            fn = hist if i < self.NH else fcst
            v = fn(i, p) if fn else None
            c = ws.cell(r, FC0 + i)
            if check:
                c.fill = fill_chk
            elif i >= self.NH:
                c.fill = fill_fc
            if v is None:
                continue
            c.value = v
            isf = isinstance(v, str) and v.startswith("=")
            c.font = (f_bold if bold else f_norm) if isf else f_in
            if isinstance(v, str) and "'Data Sources'" in v:
                c.font = f_link
            c.number_format = fmt
            if i >= self.NH and not isf and not check:
                c.fill = fill_inp
            if top:
                c.border = thin_top
        if top:
            ws.cell(r, 1).border = thin_top
        return r

    def _layout(self, ws, NC):
        R = self.R
        NH = self.NH
        segs = self.segs
        yoy = lambda key: (lambda i, p: f'=IFERROR({R(key, i)}/{R(key, i, -4)}-1,"")' if i >= 4 else None)
        hv = lambda arr: (lambda i, p: arr[i])

        # ================= DRIVERS
        self.section("OPERATING DRIVERS", NC)
        for si, sg in enumerate(segs):
            k = f"s{si}"
            nm = sg["name"]
            self.row(f"{k}_h", nm, "", None, None, bold=True)
            vol_hist = [None if v is None else v for v in sg["vol"]]
            self.row(f"{k}_vol", f"{nm}: {sg['vol_label']} growth, YoY", "%", lambda i, p, a=vol_hist: a[i],
                     lambda i, p, nm=nm: self._assume(f"seg:{nm}:volume", p, 0.02), PCT,
                     note="History: company-reported volume growth (blank where not disclosed). Forecast: input; cross-check against the industry indicators below.")
            self.row(f"{k}_px", f"{nm}: {sg['price_label']} growth, YoY", "%",
                     lambda i, p, k=k: (f'=IFERROR(IF({R(k + "_vol", i)}="","",(1+{R(k + "_g", i)})/(1+{R(k + "_vol", i)})-1),"")' if i >= 4 else None),
                     lambda i, p, nm=nm: self._assume(f"seg:{nm}:price", p, 0.02), PCT,
                     note="History implied: (1 + revenue growth) / (1 + volume growth) - 1.")
            self.row(f"{k}_rev", f"{nm}: revenue", "mm", hv(sg["revenue"]),
                     lambda i, p, k=k: f"={R(k + '_rev', i, -4)}*(1+{R(k + '_vol', i)})*(1+{R(k + '_px', i)})")
            self.row(f"{k}_g", f"  {nm}: revenue growth, YoY", "%", yoy(f"{k}_rev"), yoy(f"{k}_rev"), PCT)
            mdef = self._trail(lambda i, sg=sg: sg["gp"][i] / sg["revenue"][i])
            self.row(f"{k}_m", f"{nm}: {self.gross_label.lower()} margin", "%",
                     lambda i, p, k=k: f"=IFERROR({R(k + '_gp', i)}/{R(k + '_rev', i)},0)",
                     lambda i, p, nm=nm, mdef=mdef: self._assume(f"seg:{nm}:margin", p, round(mdef, 4)), PCT)
            self.row(f"{k}_gp", f"{nm}: {self.gross_label.lower()}", "mm", hv(sg["gp"]), lambda i, p, k=k: f"={R(k + '_rev', i)}*{R(k + '_m', i)}")
            self.blank()

        rev = self.is_["revenue"]
        gp = [r - c for r, c in zip(rev, self.is_["cost_of_revenue"])]
        self.row("h_cost", "Operating costs", "", None, None, bold=True)
        if self.headcount:
            self.row("prod", "Productivity growth (volume per employee), YoY", "%",
                     lambda i, p: f'=IFERROR(IF({R("s0_vol", i)}="","",(1+{R("s0_vol", i)})/(1+{R("hc_g", i)})-1),"")' if i >= 4 else None,
                     lambda i, p: self._assume("productivity", p, 0.03), PCT,
                     note=f"Headcount forecast = prior-year headcount x (1 + {segs[0]['name']} volume growth) / (1 + productivity growth).")
            self.row("hc", "Average headcount", "#", hv(self.headcount), lambda i, p: f"={R('hc', i, -4)}*(1+{R('s0_vol', i)})/(1+{R('prod', i)})", NUM0)
            self.row("hc_g", "  Headcount growth, YoY", "%", yoy("hc"), yoy("hc"), PCT)
            self.row("wage", "Cost per head growth, YoY", "%", lambda i, p: f'=IFERROR({R("cph", i)}/{R("cph", i, -4)}-1,"")' if i >= 4 else None,
                     lambda i, p: self._assume("wage_inflation", p, 0.03), PCT)
            hc_line = next(j for j, o in enumerate(self.opex) if o["driver"] == "headcount") if any(o["driver"] == "headcount" for o in self.opex) else None
            if hc_line is not None:
                self.row("cph", f"{self.opex[hc_line]['name']} per average head (quarterly)", "$000",
                         lambda i, p: f"=IFERROR({R(f'op{hc_line}', i)}/{R('hc', i)}*1000,0)",
                         lambda i, p: f"={R('cph', i, -4)}*(1+{R('wage', i)})", NUM)
        for j, o in enumerate(self.opex):
            if o["driver"] == "headcount":
                continue
            base = rev if o["driver"] == "pct_revenue" else gp
            lab = "revenue" if o["driver"] == "pct_revenue" else self.gross_label.lower()
            dflt = round(self._trail(lambda i, o=o, base=base: o["values"][i] / base[i]), 4)
            den = "rev" if o["driver"] == "pct_revenue" else "gp"
            self.row(f"op{j}_pct", f"{o['name']} as % of {lab}", "%", lambda i, p, j=j, den=den: f"=IFERROR({R(f'op{j}', i)}/{R(den, i)},0)",
                     lambda i, p, o=o, dflt=dflt: self._assume(f"opex:{o['name']}", p, dflt), PCT)
        da_d = round(self._trail(lambda i: self.is_["da"][i] / rev[i]), 4)
        self.row("da_pct", "Depreciation & amortization as % of revenue", "%", lambda i, p: f"=IFERROR({R('da', i)}/{R('rev', i)},0)",
                 lambda i, p: self._assume("da_pct", p, da_d), PCT2)
        sbc_d = round(self._trail(lambda i: self.cf["sbc"][i] / rev[i]), 4)
        self.row("sbc_pct", "Stock-based compensation as % of revenue", "%", lambda i, p: f"=IFERROR({R('cf_sbc', i)}/{R('rev', i)},0)",
                 lambda i, p: self._assume("sbc_pct", p, sbc_d), PCT2)
        self.blank()
        self.row("h_fin", "Financing, tax and capital return", "", None, None, bold=True)
        rate_d = round(max(0.0, self._trail(lambda i: -self.is_["interest"][i] * 4 / max(1e-9, (self.bs["short_term_debt"][i - 1] + self.bs["long_term_debt"][i - 1])) if i > 0 and (self.bs["short_term_debt"][i - 1] + self.bs["long_term_debt"][i - 1]) > 1 else None)), 4)
        self.row("rate", "Net interest & other expense / beginning total debt (annualized)", "%",
                 lambda i, p: f"=IFERROR(-{R('int', i)}*4/({R('std', i, -1)}+{R('ltd', i, -1)}),0)" if i > 0 else None,
                 lambda i, p: self._assume("interest_rate", p, rate_d or 0.05), PCT)
        tax_d = round(min(0.30, max(0.15, self._trail(lambda i: self.is_["tax"][i] / (rev[i] - self.is_["cost_of_revenue"][i] - sum(o["values"][i] for o in self.opex) - (0 if self.da_in_opex else self.is_["da"][i]) + self.is_["interest"][i])))), 4)
        self.row("tax_rate", "Effective tax rate", "%", lambda i, p: f"=IFERROR({R('tax', i)}/{R('ebt', i)},0)",
                 lambda i, p: self._assume("tax_rate", p, tax_d), PCT)
        bb_d = round(self._trail(lambda i: -self.cf["buybacks"][i]), 1)
        self.row("bb", "Share repurchases", "mm", lambda i, p: -self.cf["buybacks"][i], lambda i, p: self._assume("buybacks", p, bb_d), NUM)
        px = _num((self.spec.get("market") or {}).get("share_price")) or 0.0
        self.row("px", "Share price for repurchases", "$/sh", None, lambda i, p: self._assume("share_price", p, px), EPS,
                 note="Used to convert buyback dollars into shares retired. Set to 0 to hold the share count flat.")
        last_dps = next((v for v in reversed(self.is_["dps"]) if v), 0.0)
        self.row("dps", "Dividends declared per share", "$/sh", hv(self.is_["dps"]), lambda i, p: self._assume("dps", p, last_dps), EPS)
        self.blank()
        self.row("h_wc", "Working capital and capex", "", None, None, bold=True)
        cogs = self.is_["cost_of_revenue"]
        self.row("dso", "Days sales outstanding (receivables / revenue x 91.25)", "days", lambda i, p: f"=IFERROR({R('ar', i)}/{R('rev', i)}*91.25,0)",
                 lambda i, p: self._assume("dso", p, self._same_q_last_year(lambda j: self.bs["receivables"][j] / rev[j] * 91.25)), DAYS,
                 note="Forecast default = same fiscal quarter last year, which keeps seasonality.")
        self.row("dio", "Days inventory outstanding (inventory / cost of revenue x 91.25)", "days", lambda i, p: f"=IFERROR({R('inv', i)}/{R('cogs', i)}*91.25,0)",
                 lambda i, p: self._assume("dio", p, self._same_q_last_year(lambda j: self.bs["inventory"][j] / cogs[j] * 91.25)), DAYS)
        self.row("dpo", "Days payable outstanding (payables / cost of revenue x 91.25)", "days", lambda i, p: f"=IFERROR({R('ap', i)}/{R('cogs', i)}*91.25,0)",
                 lambda i, p: self._assume("dpo", p, self._same_q_last_year(lambda j: self.bs["payables"][j] / cogs[j] * 91.25)), DAYS)
        self.row("oca_pct", "Other current assets as % of revenue", "%", lambda i, p: f"=IFERROR({R('oca', i)}/{R('rev', i)},0)",
                 lambda i, p: self._assume("oca_pct", p, self._same_q_last_year(lambda j: self.bs["other_current_assets"][j] / rev[j])), PCT)
        self.row("ocl_pct", "Other current liabilities as % of revenue", "%", lambda i, p: f"=IFERROR({R('ocl', i)}/{R('rev', i)},0)",
                 lambda i, p: self._assume("ocl_pct", p, self._same_q_last_year(lambda j: self.bs["other_current_liabilities"][j] / rev[j])), PCT)
        cap_d = round(self._trail(lambda i: -self.cf["capex"][i] / rev[i]), 4)
        self.row("capex_pct", "Capital expenditures as % of revenue", "%", lambda i, p: f"=IFERROR(-{R('cf_cap', i)}/{R('rev', i)},0)",
                 lambda i, p: self._assume("capex_pct", p, cap_d), PCT2)
        mc_d = round(0.5 * self.bs["cash"][-1], 1)
        self.row("mincash", "Minimum cash balance (revolver draws below this)", "mm", None, lambda i, p: self._assume("min_cash", p, mc_d), NUM)
        self.blank()

        # ================= INCOME STATEMENT
        self.section("INCOME STATEMENT", NC)
        segsum = lambda key: (lambda i, p: "=" + "+".join(R(f"s{si}_{key}", i) for si in range(len(segs))))
        self.row("rev", "Total revenues", "mm", hv(rev), segsum("rev"), bold=True)
        self.row("rev_g", "  Revenue growth, YoY", "%", yoy("rev"), yoy("rev"), PCT)
        self.row("cogs", self.cogs_label, "mm", hv(cogs), lambda i, p: f"={R('rev', i)}-(" + "+".join(R(f"s{si}_gp", i) for si in range(len(segs))) + ")")
        f_gp = lambda i, p: f"={R('rev', i)}-{R('cogs', i)}"
        self.row("gp", self.gross_label, "mm", f_gp, f_gp, bold=True, top=True)
        f_gpm = lambda i, p: f"=IFERROR({R('gp', i)}/{R('rev', i)},0)"
        self.row("gpm", f"  {self.gross_label} margin", "%", f_gpm, f_gpm, PCT)
        for j, o in enumerate(self.opex):
            if o["driver"] == "headcount":
                fc = lambda i, p: f"={R('hc', i)}*{R('cph', i)}/1000"
            else:
                den = "rev" if o["driver"] == "pct_revenue" else "gp"
                fc = lambda i, p, j=j, den=den: f"={R(den, i)}*{R(f'op{j}_pct', i)}"
            self.row(f"op{j}", o["name"], "mm", lambda i, p, o=o: o["values"][i], fc)
        if not self.da_in_opex:
            self.row("da_exp", "Depreciation & amortization", "mm", lambda i, p: f"={R('da', i)}", lambda i, p: f"={R('da', i)}")
        opk = [f"op{j}" for j in range(len(self.opex))] + ([] if self.da_in_opex else ["da_exp"])
        f_ebit = lambda i, p: f"={R('gp', i)}-(" + "+".join(R(k, i) for k in opk) + ")"
        self.row("ebit", "Operating income", "mm", f_ebit, f_ebit, bold=True, top=True)
        f_om = lambda i, p: f"=IFERROR({R('ebit', i)}/{R('rev', i)},0)"
        self.row("ebit_m", "  Operating margin", "%", f_om, f_om, PCT)
        self.row("int", "Interest and other income (expense), net", "mm", hv(self.is_["interest"]),
                 lambda i, p: f"=-{R('rate', i)}/4*({R('std', i, -1)}+{R('ltd', i, -1)})")
        f_ebt = lambda i, p: f"={R('ebit', i)}+{R('int', i)}"
        self.row("ebt", "Pre-tax income", "mm", f_ebt, f_ebt, top=True)
        self.row("tax", "Income tax expense", "mm", hv(self.is_["tax"]), lambda i, p: f"={R('ebt', i)}*{R('tax_rate', i)}")
        f_ni = lambda i, p: f"={R('ebt', i)}-{R('tax', i)}"
        self.row("ni", "Net income", "mm", f_ni, f_ni, bold=True, top=True)
        self.row("sh", "Diluted weighted average shares", "mm", hv(self.is_["shares"]),
                 lambda i, p: f"=MAX({R('sh', i, -1)}-IFERROR({R('bb', i)}/{R('px', i)},0),0)")
        f_eps = lambda i, p: f"=IFERROR({R('ni', i)}/{R('sh', i)},0)"
        self.row("eps", "Diluted EPS", "$/sh", f_eps, f_eps, EPS, bold=True)
        if any(v is not None for v in self.is_["eps_rep"]):
            self.row("eps_rep", "  Reported diluted EPS (memo)", "$/sh", hv(self.is_["eps_rep"]), None, EPS)
        self.blank()
        self.row("h_memo", "Memo", "", None, None, bold=True)
        self.row("da", "Depreciation & amortization", "mm", hv(self.is_["da"]), lambda i, p: f"={R('rev', i)}*{R('da_pct', i)}",
                 note="Included within operating expenses above." if self.da_in_opex else "Shown as its own expense line above.")
        f_ebitda = lambda i, p: f"={R('ebit', i)}+{R('da', i)}"
        self.row("ebitda", "EBITDA", "mm", f_ebitda, f_ebitda)
        self.row("adj", "Non-GAAP adjustments to operating income (pre-tax)", "mm", hv(self.is_["adj"]), lambda i, p: 0.0)
        f_ae = lambda i, p: f"={R('ebit', i)}+{R('adj', i)}"
        self.row("adj_ebit", "Adjusted operating income", "mm", f_ae, f_ae)
        if len(segs) > 1:
            self.row("seg_chk", "Check: segment revenue - total revenue (should be 0)", "mm",
                     lambda i, p: "=ROUND(" + "+".join(R(f"s{si}_rev", i) for si in range(len(segs))) + f"-{R('rev', i)},3)", None, NUM, check=True)
        if any(v is not None for v in self.is_["ni_rep"]):
            self.row("ni_chk", "Check: model net income - reported (should be 0)", "mm",
                     lambda i, p: (f"=ROUND({R('ni', i)}-{self.is_['ni_rep'][i]},1)" if self.is_["ni_rep"][i] is not None else None), None, NUM, check=True)
        self.blank()

        # ================= BALANCE SHEET
        self.section("BALANCE SHEET (period end)", NC)
        b = self.bs
        self.row("cash", "Cash and equivalents", "mm", hv(b["cash"]), lambda i, p: f"={R('cash_pre', i)}+{R('std', i)}-{R('std', i, -1)}")
        self.row("ar", "Receivables", "mm", hv(b["receivables"]), lambda i, p: f"={R('dso', i)}/91.25*{R('rev', i)}")
        self.row("inv", "Inventories", "mm", hv(b["inventory"]), lambda i, p: f"={R('dio', i)}/91.25*{R('cogs', i)}")
        self.row("oca", "Other current assets", "mm", hv(b["other_current_assets"]), lambda i, p: f"={R('oca_pct', i)}*{R('rev', i)}")
        f_tca = lambda i, p: f"=SUM({R('cash', i)}:{R('oca', i)})"
        self.row("tca", "Total current assets", "mm", f_tca, f_tca, top=True)
        self.row("ppe", "PP&E and intangibles, net", "mm", hv(b["ppe_intangibles"]), lambda i, p: f"={R('ppe', i, -1)}-{R('cf_cap', i)}-{R('da', i)}")
        self.row("gw", "Goodwill", "mm", hv(b["goodwill"]), lambda i, p: f"={R('gw', i, -1)}-{R('cf_acq', i)}")
        self.row("onca", "Other non-current assets", "mm", hv(b["other_noncurrent_assets"]), lambda i, p: f"={R('onca', i, -1)}-{R('cf_oi', i)}")
        f_ta = lambda i, p: f"={R('tca', i)}+SUM({R('ppe', i)}:{R('onca', i)})"
        self.row("ta", "Total assets", "mm", f_ta, f_ta, bold=True, top=True)
        self.blank()
        self.row("ap", "Accounts payable", "mm", hv(b["payables"]), lambda i, p: f"={R('dpo', i)}/91.25*{R('cogs', i)}")
        self.row("ocl", "Other current liabilities", "mm", hv(b["other_current_liabilities"]), lambda i, p: f"={R('ocl_pct', i)}*{R('rev', i)}")
        self.row("std", "Short-term debt / revolver", "mm", hv(b["short_term_debt"]), lambda i, p: f"=MAX(0,{R('std', i, -1)}+{R('mincash', i)}-{R('cash_pre', i)})",
                 note="Forecast plug: draws when cash would fall below the minimum and is repaid first from excess cash.")
        f_tcl = lambda i, p: f"=SUM({R('ap', i)}:{R('std', i)})"
        self.row("tcl", "Total current liabilities", "mm", f_tcl, f_tcl, top=True)
        self.row("ltd", "Long-term debt", "mm", hv(b["long_term_debt"]), lambda i, p: f"={R('ltd', i, -1)}", note="Held flat in the forecast; edit to model maturities or issuance.")
        self.row("oncl", "Other non-current liabilities", "mm", hv(b["other_noncurrent_liabilities"]), lambda i, p: f"={R('oncl', i, -1)}+{R('cf_nc', i)}")
        f_tl = lambda i, p: f"={R('tcl', i)}+{R('ltd', i)}+{R('oncl', i)}"
        self.row("tl", "Total liabilities", "mm", f_tl, f_tl, top=True)
        self.row("eq", "Shareholders' equity", "mm", hv(b["equity"]),
                 lambda i, p: f"={R('eq', i, -1)}+{R('ni', i)}+{R('cf_sbc', i)}+{R('cf_bb', i)}+{R('cf_div', i)}+{R('cf_of', i)}+{R('cf_fx', i)}")
        f_tle = lambda i, p: f"={R('tl', i)}+{R('eq', i)}"
        self.row("tle", "Total liabilities & equity", "mm", f_tle, f_tle, bold=True, top=True)
        f_bc = lambda i, p: f"=ROUND({R('ta', i)}-{R('tle', i)},3)"
        self.row("bs_chk", "Check: assets - liabilities & equity (should be 0)", "mm", f_bc, f_bc, NUM, check=True)
        self.blank()

        # ================= CASH FLOW
        self.section("CASH FLOW STATEMENT", NC)
        c = self.cf
        self.row("cf_ni", "Net income", "mm", lambda i, p: f"={R('ni', i)}", lambda i, p: f"={R('ni', i)}")
        self.row("cf_da", "Depreciation & amortization", "mm", lambda i, p: f"={R('da', i)}", lambda i, p: f"={R('da', i)}")
        self.row("cf_sbc", "Stock-based compensation", "mm", hv(c["sbc"]), lambda i, p: f"={R('rev', i)}*{R('sbc_pct', i)}")
        self.row("cf_nc", "Other non-cash items & reconciling items", "mm",
                 lambda i, p: c["other_noncash"][i] + self.plugs["cfo"][i], lambda i, p: 0.0,
                 note="History includes any difference to reported operating cash flow so the statement ties to the filing.")
        self.row("cf_ar", "Change in receivables", "mm", hv(c["change_receivables"]), lambda i, p: f"={R('ar', i, -1)}-{R('ar', i)}")
        self.row("cf_inv", "Change in inventories", "mm", hv(c["change_inventory"]), lambda i, p: f"={R('inv', i, -1)}-{R('inv', i)}")
        self.row("cf_ap", "Change in payables", "mm", hv(c["change_payables"]), lambda i, p: f"={R('ap', i)}-{R('ap', i, -1)}")
        self.row("cf_ow", "Change in other working capital", "mm", hv(c["change_other_wc"]),
                 lambda i, p: f"={R('oca', i, -1)}-{R('oca', i)}+{R('ocl', i)}-{R('ocl', i, -1)}")
        f_cfo = lambda i, p: f"=SUM({R('cf_ni', i)}:{R('cf_ow', i)})"
        self.row("cfo", "Net cash from operating activities", "mm", f_cfo, f_cfo, bold=True, top=True)
        self.row("cf_cap", "Capital expenditures", "mm", hv(c["capex"]), lambda i, p: f"=-{R('capex_pct', i)}*{R('rev', i)}")
        self.row("cf_acq", "Acquisitions, net of cash acquired", "mm", hv(c["acquisitions"]), lambda i, p: self._assume("acquisitions", p, 0.0))
        self.row("cf_oi", "Other investing (incl. reconciling items)", "mm", lambda i, p: c["other_investing"][i] + self.plugs["cfi"][i], lambda i, p: 0.0)
        f_cfi = lambda i, p: f"=SUM({R('cf_cap', i)}:{R('cf_oi', i)})"
        self.row("cfi", "Net cash from investing activities", "mm", f_cfi, f_cfi, bold=True, top=True)
        self.row("cf_bb", "Repurchases of common stock", "mm", lambda i, p: f"=-{R('bb', i)}", lambda i, p: f"=-{R('bb', i)}")
        self.row("cf_div", "Dividends paid", "mm", hv(c["dividends"]), lambda i, p: f"=-{R('dps', i)}*{R('sh', i)}",
                 note="Forecast: dividends per share x diluted shares.")
        self.row("cf_debt", "Debt issued (repaid), net", "mm", hv(c["debt_net"]),
                 lambda i, p: f"={R('std', i)}-{R('std', i, -1)}+{R('ltd', i)}-{R('ltd', i, -1)}")
        self.row("cf_of", "Other financing (incl. reconciling items)", "mm", lambda i, p: c["other_financing"][i] + self.plugs["cff"][i], lambda i, p: 0.0)
        f_cff = lambda i, p: f"=SUM({R('cf_bb', i)}:{R('cf_of', i)})"
        self.row("cff", "Net cash from financing activities", "mm", f_cff, f_cff, bold=True, top=True)
        self.row("cf_fx", "Effect of exchange rates", "mm", hv(c["fx"]), lambda i, p: 0.0)
        f_net = lambda i, p: f"={R('cfo', i)}+{R('cfi', i)}+{R('cff', i)}+{R('cf_fx', i)}"
        self.row("cf_net", "Net change in cash", "mm", f_net, f_net, bold=True, top=True)
        f_fcf = lambda i, p: f"={R('cfo', i)}+{R('cf_cap', i)}"
        self.row("fcf", "Free cash flow (operating cash flow + capex)", "mm", f_fcf, f_fcf)
        self.row("cash_pre", "Memo: cash before revolver draw / (repayment)", "mm", None,
                 lambda i, p: f"={R('cash', i, -1)}+{R('cfo', i)}+{R('cfi', i)}+{R('cf_bb', i)}+{R('cf_div', i)}+{R('cf_of', i)}+{R('cf_fx', i)}+{R('ltd', i)}-{R('ltd', i, -1)}")
        cfo_rep = c["cfo_reported"]
        self.row("cfo_chk", "Check: reported operating cash flow - model (should be 0)", "mm", lambda i, p: f"=ROUND({cfo_rep[i]}-{R('cfo', i)},1)", None, NUM, check=True)
        open_cash = self.opening.get("cash")
        self.row("cash_chk", "Check: cash flow vs. balance sheet cash change", "mm",
                 lambda i, p: (f"=ROUND({R('cf_net', i)}-({R('cash', i)}-{R('cash', i, -1)}),1)" if i > 0 else
                               (f"=ROUND({R('cf_net', i)}-({R('cash', i)}-{open_cash}),1)" if open_cash is not None else None)),
                 lambda i, p: f"=ROUND({R('cf_net', i)}-({R('cash', i)}-{R('cash', i, -1)}),1)", NUM, check=True,
                 note="Historical differences usually come from restricted cash or cash held for sale; forecast must be 0.")
        if self._pass == 2:
            for i in range(1, self.NH):
                dlt = c["cfo_reported"][i] + c["cfi_reported"][i] + c["cff_reported"][i] + c["fx"][i] - (b["cash"][i] - b["cash"][i - 1])
                if abs(dlt) > max(5.0, 0.02 * abs(b["cash"][i])):
                    self.rep.warnings.append(f"{self.hist[i]['label']}: reported cash flows differ from the balance sheet cash change by {dlt:,.1f} (restricted cash, held-for-sale or a data issue)")
                    break

        # ================= industry memo
        inds = self.spec.get("industry_series") or []
        if inds:
            self.blank()
            self.section("INDUSTRY INDICATORS, YoY (linked from Data Sources)", NC)
            self.ind_links = []
            for n_, s_ in enumerate(inds):
                r = self.row(f"ind{n_}", s_.get("name", f"Series {n_ + 1}"), "%", None, None)
                self.ind_links.append(r)
        self.qrows = dict(self.rows)

    # ----- annual
    def _qcols(self, y):
        idx = [i for i, p in enumerate(self.periods) if p["fy"] == y]
        return L(FC0 + idx[0]), L(FC0 + idx[-1])

    def _annual(self, ws):
        self.am = ws
        QS = "'Quarterly Model'"
        ws["A1"] = f"{self.spec.get('company', self.spec.get('ticker'))} - Annual Model (rolls up from Quarterly Model)"; ws["A1"].font = f_title
        ws["A2"] = "Income statement and cash flow = sum of four fiscal quarters; balance sheet = fiscal year-end value."; ws["A2"].font = Font(name=FONT, size=9, italic=True)
        last_hist_fy = self.hist[-1]["fy"] if self.hist[-1]["fq"] == 4 else self.hist[-1]["fy"] - 1
        for j, y in enumerate(self.years):
            c = FC0 + j
            ws.cell(4, c, f"FY{y}{'A' if y <= last_hist_fy else 'E'}").font = f_hdr
            q4 = [p for p in self.periods if p["fy"] == y][-1]
            d = ws.cell(5, c, q4["end"]); d.number_format = "m/d/yyyy"
            ws.cell(6, c, "Actual" if y <= last_hist_fy else ("Actual + Forecast" if any(p["fy"] == y for p in self.hist) else "Forecast")).font = f_small
            for r in (4, 5, 6):
                ws.cell(r, c).alignment = Alignment(horizontal="right")
                if y > last_hist_fy:
                    ws.cell(r, c).fill = fill_fc
        ws.cell(4, 1, "Fiscal year").font = f_hdr
        ws.cell(5, 1, "Fiscal year end").font = f_hdr
        self.arows = {}
        cur = [7]
        q = self.qrows

        def sec(t):
            r = cur[0]; cur[0] += 1
            ws.cell(r, 1, t).font = f_sec
            for k in range(1, FC0 + len(self.years)):
                ws.cell(r, k).fill = fill_sec

        def arow(key, label, kind, fmt=NUM, bold=False, top=False, formula=None, src=None):
            r = cur[0]; cur[0] += 1
            self.arows[key] = r
            ws.cell(r, 1, label).font = f_bold if bold else f_norm
            for j, y in enumerate(self.years):
                c = ws.cell(r, FC0 + j)
                a, bcol = self._qcols(y)
                qr = q.get(src or key)
                if kind == "f":
                    c.value = formula(j); c.font = f_bold if bold else f_norm
                elif kind == "sum":
                    c.value = f"=SUM({QS}!{a}{qr}:{bcol}{qr})"
                elif kind == "end":
                    c.value = f"={QS}!{bcol}{qr}"
                elif kind == "avg":
                    c.value = f"=AVERAGE({QS}!{a}{qr}:{bcol}{qr})"
                if kind != "f":
                    c.font = Font(name=FONT, size=10, bold=bold, color="008000")
                c.number_format = fmt
                if top:
                    c.border = thin_top
                if y > last_hist_fy:
                    c.fill = fill_fc

        AR = lambda k, j, off=0: f"{L(FC0 + j + off)}{self.arows[k]}"
        g = lambda k: (lambda j: f'=IFERROR({AR(k, j)}/{AR(k, j, -1)}-1,"")' if j else "")
        sec("SEGMENTS")
        for si, sg in enumerate(self.segs):
            arow(f"s{si}_rev", f"{sg['name']}: revenue", "sum")
            arow(f"s{si}_g", "  growth", "f", PCT, formula=g(f"s{si}_rev"))
            arow(f"s{si}_gp", f"{sg['name']}: {self.gross_label.lower()}", "sum")
        if self.headcount:
            arow("hc", "Average headcount", "avg", NUM0)
        cur[0] += 1
        sec("INCOME STATEMENT")
        arow("rev", "Total revenues", "sum", bold=True)
        arow("rev_g", "  Revenue growth", "f", PCT, formula=g("rev"))
        arow("cogs", self.cogs_label, "sum")
        arow("gp", self.gross_label, "sum", bold=True, top=True)
        arow("gpm", f"  {self.gross_label} margin", "f", PCT, formula=lambda j: f"=IFERROR({AR('gp', j)}/{AR('rev', j)},0)")
        for jx, o in enumerate(self.opex):
            arow(f"op{jx}", o["name"], "sum")
        if not self.da_in_opex:
            arow("da_exp", "Depreciation & amortization", "sum")
        arow("ebit", "Operating income", "sum", bold=True, top=True)
        arow("ebit_m", "  Operating margin", "f", PCT, formula=lambda j: f"=IFERROR({AR('ebit', j)}/{AR('rev', j)},0)")
        arow("int", "Interest and other, net", "sum")
        arow("ebt", "Pre-tax income", "sum", top=True)
        arow("tax", "Income tax expense", "sum")
        arow("ni", "Net income", "sum", bold=True, top=True)
        arow("sh", "Diluted shares (average)", "avg")
        arow("eps", "Diluted EPS (sum of quarters)", "sum", EPS, bold=True)
        arow("dps", "Dividends per share", "sum", EPS)
        arow("da", "Depreciation & amortization", "sum")
        arow("ebitda", "EBITDA", "sum")
        arow("adj_ebit", "Adjusted operating income", "sum")
        cur[0] += 1
        sec("BALANCE SHEET (fiscal year end)")
        for k, lab in [("cash", "Cash and equivalents"), ("ar", "Receivables"), ("inv", "Inventories"), ("oca", "Other current assets"), ("tca", "Total current assets"),
                       ("ppe", "PP&E and intangibles, net"), ("gw", "Goodwill"), ("onca", "Other non-current assets"), ("ta", "Total assets"),
                       ("ap", "Accounts payable"), ("ocl", "Other current liabilities"), ("std", "Short-term debt / revolver"), ("tcl", "Total current liabilities"),
                       ("ltd", "Long-term debt"), ("oncl", "Other non-current liabilities"), ("tl", "Total liabilities"), ("eq", "Shareholders' equity"),
                       ("tle", "Total liabilities & equity")]:
            arow(k, lab, "end", bold=k in ("ta", "tle"), top=k in ("tca", "ta", "tcl", "tl", "tle"))
        arow("bs_chk", "Check: assets - liabilities & equity", "f", formula=lambda j: f"=ROUND({AR('ta', j)}-{AR('tle', j)},3)")
        cur[0] += 1
        sec("CASH FLOW STATEMENT")
        for k, lab in [("cf_ni", "Net income"), ("cf_da", "Depreciation & amortization"), ("cf_sbc", "Stock-based compensation"), ("cf_nc", "Other non-cash items"),
                       ("cf_ar", "Change in receivables"), ("cf_inv", "Change in inventories"), ("cf_ap", "Change in payables"), ("cf_ow", "Change in other working capital"),
                       ("cfo", "Net cash from operating activities"), ("cf_cap", "Capital expenditures"), ("cf_acq", "Acquisitions"), ("cf_oi", "Other investing"),
                       ("cfi", "Net cash from investing activities"), ("cf_bb", "Share repurchases"), ("cf_div", "Dividends paid"), ("cf_debt", "Debt, net"),
                       ("cf_of", "Other financing"), ("cff", "Net cash from financing activities"), ("cf_fx", "FX effect"), ("cf_net", "Net change in cash"),
                       ("fcf", "Free cash flow")]:
            arow(k, lab, "sum", bold=k in ("cfo", "cfi", "cff", "cf_net"), top=k in ("cfo", "cfi", "cff", "cf_net"))
        ws.column_dimensions["A"].width = 46
        ws.column_dimensions["B"].width = 3
        for j in range(len(self.years)):
            ws.column_dimensions[L(FC0 + j)].width = 12
        ws.freeze_panes = "C7"

    # ----- key metrics
    def _metrics(self, ws):
        AMS = "'Annual Model'"
        ws["A1"] = f"{self.spec.get('company', self.spec.get('ticker'))} - Key Metrics (fiscal years)"; ws["A1"].font = f_title
        ws["A2"] = "Averages use beginning and ending fiscal-year balances (first year uses year-end only). Days metrics use a 365-day year."
        ws["A2"].font = Font(name=FONT, size=9, italic=True)
        last_hist_fy = self.hist[-1]["fy"] if self.hist[-1]["fq"] == 4 else self.hist[-1]["fy"] - 1
        for j, y in enumerate(self.years):
            c = ws.cell(4, FC0 + j, f"FY{y}{'A' if y <= last_hist_fy else 'E'}"); c.font = f_hdr; c.alignment = Alignment(horizontal="right")
            if y > last_hist_fy:
                c.fill = fill_fc
        A = lambda k, j, off=0: f"{AMS}!{L(FC0 + j + off)}{self.arows[k]}"
        avg = lambda k, j: A(k, j) if j == 0 else f"AVERAGE({A(k, j, -1)},{A(k, j)})"
        cur = [6]
        rows = {}

        def sec(t):
            r = cur[0]; cur[0] += 1
            ws.cell(r, 1, t).font = f_sec
            for k in range(1, FC0 + len(self.years)):
                ws.cell(r, k).fill = fill_sec

        def m(key, label, unit, f, fmt, note=None):
            r = cur[0]; cur[0] += 1
            rows[key] = r
            ws.cell(r, 1, label).font = f_norm
            ws.cell(r, 2, unit).font = f_small
            if note:
                ws.cell(r, 1).comment = Comment(note, "Model")
            for j, y in enumerate(self.years):
                c = ws.cell(r, FC0 + j, f(j)); c.number_format = fmt; c.font = f_norm
                if y > last_hist_fy:
                    c.fill = fill_fc
        K = lambda k, j: f"{L(FC0 + j)}{rows[k]}"
        ic = lambda j, off=0: f"({A('std', j, off)}+{A('ltd', j, off)}+{A('eq', j, off)}-{A('cash', j, off)})"
        sec("RETURNS")
        m("roe", "Return on equity (RoE)", "%", lambda j: f"=IFERROR({A('ni', j)}/{avg('eq', j)},0)", PCT)
        m("roa", "Return on assets (RoA)", "%", lambda j: f"=IFERROR({A('ni', j)}/{avg('ta', j)},0)", PCT)
        m("roic", "Return on invested capital (RoIC)", "%",
          lambda j: f"=IFERROR({A('ebit', j)}*(1-{A('tax', j)}/{A('ebt', j)})/" + (ic(j) if j == 0 else f"AVERAGE({ic(j, -1)},{ic(j)})") + ",0)", PCT,
          note="NOPAT = operating income x (1 - effective tax rate); invested capital = total debt + equity - cash.")
        sec("WORKING CAPITAL AND CASH CYCLE")
        m("dso", "Days sales outstanding (DSO)", "days", lambda j: f"=IFERROR({A('ar', j)}/{A('rev', j)}*365,0)", DAYS)
        m("dio", "Days inventory outstanding (DIO)", "days", lambda j: f"=IFERROR({A('inv', j)}/{A('cogs', j)}*365,0)", DAYS)
        m("dpo", "Days payable outstanding (DPO)", "days", lambda j: f"=IFERROR({A('ap', j)}/{A('cogs', j)}*365,0)", DAYS)
        m("turns", "Inventory turns (cost of revenue / average inventory)", "x",
          lambda j: f'=IFERROR(IF({avg("inv", j)}=0,"n/m",{A("cogs", j)}/{avg("inv", j)}),"n/m")', MULT,
          note="Shows n/m when the company carries no inventory.")
        m("ccc", "Cash conversion cycle (DSO + DIO - DPO)", "days", lambda j: f"={K('dso', j)}+{K('dio', j)}-{K('dpo', j)}", DAYS)
        sec("LEVERAGE AND COVERAGE")
        m("debt", "Total debt", "mm", lambda j: f"={A('std', j)}+{A('ltd', j)}", NUM)
        m("nd", "Net debt", "mm", lambda j: f"={A('std', j)}+{A('ltd', j)}-{A('cash', j)}", NUM)
        m("ebitda", "EBITDA", "mm", lambda j: f"={A('ebitda', j)}", NUM)
        m("ndx", "Net debt / EBITDA", "x", lambda j: f"=IFERROR({K('nd', j)}/{K('ebitda', j)},0)", MULT)
        m("de", "Total debt / equity", "x", lambda j: f"=IFERROR({K('debt', j)}/{A('eq', j)},0)", MULT)
        m("cov", "Interest coverage (operating income / net interest)", "x", lambda j: f'=IFERROR(IF({A("int", j)}>=0,"n/m",{A("ebit", j)}/-{A("int", j)}),"n/m")', MULT)
        sec("PROFITABILITY")
        m("rg", "Revenue growth", "%", lambda j: f'=IFERROR({A("rev", j)}/{A("rev", j, -1)}-1,"")' if j else "", PCT)
        m("gpm", f"{self.gross_label} margin", "%", lambda j: f"=IFERROR({A('gp', j)}/{A('rev', j)},0)", PCT)
        m("om", "Operating margin", "%", lambda j: f"=IFERROR({A('ebit', j)}/{A('rev', j)},0)", PCT)
        m("aom", "Adjusted operating margin", "%", lambda j: f"=IFERROR({A('adj_ebit', j)}/{A('rev', j)},0)", PCT)
        m("om_gp", f"Operating income / {self.gross_label.lower()}", "%", lambda j: f"=IFERROR({A('ebit', j)}/{A('gp', j)},0)", PCT)
        m("em", "EBITDA margin", "%", lambda j: f"=IFERROR({A('ebitda', j)}/{A('rev', j)},0)", PCT)
        m("nm", "Net margin", "%", lambda j: f"=IFERROR({A('ni', j)}/{A('rev', j)},0)", PCT)
        if self.headcount:
            m("rpe", "Revenue per average employee", "$000", lambda j: f"=IFERROR({A('rev', j)}/{A('hc', j)}*1000,0)", NUM)
            m("gpe", f"{self.gross_label} per average employee", "$000", lambda j: f"=IFERROR({A('gp', j)}/{A('hc', j)}*1000,0)", NUM)
        sec("CASH FLOW AND CAPITAL RETURN")
        m("fcf", "Free cash flow", "mm", lambda j: f"={A('fcf', j)}", NUM)
        m("fcfc", "FCF conversion (FCF / net income)", "%", lambda j: f"=IFERROR({A('fcf', j)}/{A('ni', j)},0)", PCT)
        m("capex", "Capex as % of revenue", "%", lambda j: f"=IFERROR(-{A('cf_cap', j)}/{A('rev', j)},0)", PCT2)
        m("payout", "Shareholder payout ((dividends + buybacks) / net income)", "%", lambda j: f"=IFERROR(-({A('cf_div', j)}+{A('cf_bb', j)})/{A('ni', j)},0)", PCT)
        ws.column_dimensions["A"].width = 52
        ws.column_dimensions["B"].width = 7
        for j in range(len(self.years)):
            ws.column_dimensions[L(FC0 + j)].width = 12
        ws.freeze_panes = "C5"

    # ----- competitors
    def _competitors(self, ws):
        comps = self.spec.get("competitors") or []
        ws["A1"] = "Competitor Benchmarking and Relative Market Share"; ws["A1"].font = f_title
        ws["A2"] = (self.spec.get("competitor_basis") or "USD millions. FY = latest fiscal year; LTM = last twelve months. "
                    "Share = each company's revenue or operating profit as a % of the peer set shown.")
        ws["A2"].font = Font(name=FONT, size=9, italic=True)
        hdr = ["Ticker", "Company", "Business", "Revenue (FY)", "Gross metric (FY)", "Gross metric basis", "Operating income (FY)", "EBITDA (FY)",
               "Net income (FY)", "Revenue (LTM)", "Operating income (LTM)", "Net income (LTM)", "Market cap", "Total debt", "Cash", "Enterprise value",
               "Gross metric / revenue", "Operating margin (FY)", "EV / EBITDA", "P / E (LTM)", "Revenue share (FY)", "Operating profit share (FY)",
               "Revenue share (LTM)", "Operating profit share (LTM)"]
        fields = {"Revenue (FY)": "revenue", "Gross metric (FY)": "gross_profit", "Operating income (FY)": "operating_income", "EBITDA (FY)": "ebitda",
                  "Net income (FY)": "net_income", "Revenue (LTM)": "ltm_revenue", "Operating income (LTM)": "ltm_operating_income",
                  "Net income (LTM)": "ltm_net_income", "Market cap": "market_cap", "Total debt": "total_debt", "Cash": "cash"}
        HR = 4
        Cx = {h: L(1 + k) for k, h in enumerate(hdr)}
        for k, h in enumerate(hdr):
            c = ws.cell(HR, 1 + k, h); c.font = f_hdr; c.fill = fill_th; c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
        ws.row_dimensions[HR].height = 45
        if not comps:
            ws.cell(HR + 1, 1, "No competitor data was provided.").font = f_small
            return
        r0 = HR + 1
        r1 = r0 + len(comps) - 1
        for n, cp in enumerate(comps):
            r = r0 + n
            ws.cell(r, 1, cp.get("ticker")).font = f_bold
            ws.cell(r, 2, cp.get("name")).font = f_norm
            ws.cell(r, 3, cp.get("business")).font = f_norm
            for h, fld in fields.items():
                v = _num(cp.get(fld))
                if v is not None:
                    c = ws[f"{Cx[h]}{r}"]; c.value = round(v, 1); c.font = f_in; c.number_format = NUM
            ws[f"{Cx['Gross metric basis']}{r}"] = cp.get("gross_metric_label") or "n/a"
            ws[f"{Cx['Gross metric basis']}{r}"].font = Font(name=FONT, size=9)
            g = lambda h: f"{Cx[h]}{r}"

            def put(h, f, fmt):
                c = ws[g(h)]; c.value = f; c.number_format = fmt; c.font = f_norm
            put("Enterprise value", f'=IF(OR({g("Market cap")}="",{g("Total debt")}=""),"",{g("Market cap")}+{g("Total debt")}-N({g("Cash")}))', NUM)
            put("Gross metric / revenue", f'=IFERROR({g("Gross metric (FY)")}/{g("Revenue (FY)")},"n/a")', PCT)
            put("Operating margin (FY)", f'=IFERROR({g("Operating income (FY)")}/{g("Revenue (FY)")},"n/a")', PCT)
            put("EV / EBITDA", f'=IFERROR(IF({g("EBITDA (FY)")}<=0,"n/m",{g("Enterprise value")}/{g("EBITDA (FY)")}),"n/a")', MULT)
            put("P / E (LTM)", f'=IFERROR(IF({g("Net income (LTM)")}<=0,"n/m",{g("Market cap")}/{g("Net income (LTM)")}),"n/a")', MULT)
            for share, src, kind in [("Revenue share (FY)", "Revenue (FY)", "rev"), ("Operating profit share (FY)", "Operating income (FY)", "op"),
                                     ("Revenue share (LTM)", "Revenue (LTM)", "rev"), ("Operating profit share (LTM)", "Operating income (LTM)", "op")]:
                col = Cx[src]
                if kind == "rev":
                    put(share, f'=IF(ISNUMBER({col}{r}),{col}{r}/SUM(${col}${r0}:${col}${r1}),"n/a")', PCT)
                else:
                    put(share, f'=IF(ISNUMBER({col}{r}),MAX({col}{r},0)/SUMIF(${col}${r0}:${col}${r1},">0"),"n/a")', PCT)
            if (cp.get("ticker") or "").upper() == (self.spec.get("ticker") or "").upper():
                for k in range(1, len(hdr) + 1):
                    ws.cell(r, k).fill = fill_inp
        tr = r1 + 1
        ws.cell(tr, 2, "Peer set total").font = f_bold
        for h in ["Revenue (FY)", "Operating income (FY)", "Revenue (LTM)", "Operating income (LTM)", "Market cap"]:
            c = ws[f"{Cx[h]}{tr}"]; c.value = f"=SUM({Cx[h]}{r0}:{Cx[h]}{r1})"; c.font = f_bold; c.number_format = NUM; c.border = thin_top
        nr = tr + 2
        notes = ["Notes"] + [str(x) for x in (self.spec.get("competitor_notes") or [])] + [
            "Operating profit share counts only companies with positive operating income.",
            "Clickable source links for each company are on the Data Sources tab."]
        for k, t in enumerate(notes):
            ws.cell(nr + k, 1, t).font = f_bold if k == 0 else Font(name=FONT, size=9)
        for k, w in enumerate([10, 32, 24] + [12] * (len(hdr) - 3)):
            ws.column_dimensions[L(1 + k)].width = w
        ws.column_dimensions[Cx["Gross metric basis"]].width = 20
        ws.freeze_panes = "D5"

    # ----- data sources
    def _sources(self, ws):
        s = self.spec
        ws["A1"] = "Data Sources"; ws["A1"].font = f_title
        ws["A2"] = "Every source link is a clickable cell. Values on the model tabs marked blue were taken from these sources."
        ws["A2"].font = Font(name=FONT, size=9, italic=True)
        r = 4

        def head(cols):
            nonlocal r
            for k, h in enumerate(cols):
                c = ws.cell(r, 1 + k, h); c.font = f_hdr; c.fill = fill_th
            r += 1

        def link(rr, cc, u):
            c = ws.cell(rr, cc, u); c.hyperlink = u; c.font = f_url

        ws.cell(r, 1, "Company filings and documents").font = f_bold; r += 1
        head(["Document", "Category", "Link"])
        for src in s.get("sources") or []:
            ws.cell(r, 1, src.get("label")).font = Font(name=FONT, size=9)
            ws.cell(r, 2, src.get("category", "")).font = Font(name=FONT, size=9)
            if src.get("url"):
                link(r, 3, src["url"])
            r += 1
        r += 2
        inds = s.get("industry_series") or []
        if inds:
            hist_q = [f"{p['fy']}Q{p['fq']}" for p in self.periods]
            qcols = sorted({k for sr in inds for k in (sr.get("values") or {})}, key=lambda x: (int(x[:4]), int(x[-1])))
            ws.cell(r, 1, "Industry and government indicators (levels)").font = f_bold; r += 1
            hdr = ["Series", "Publisher", "Frequency", "Unit"] + [f"Q{q[-1]} FY{q[:4]}" for q in qcols]
            head(hdr)
            lvl = {}
            for n, sr in enumerate(inds):
                lvl[n] = r
                for k, f in enumerate(["name", "publisher", "frequency", "unit"]):
                    ws.cell(r, 1 + k, sr.get(f)).font = Font(name=FONT, size=9)
                for k, q in enumerate(qcols):
                    v = _num((sr.get("values") or {}).get(q))
                    if v is not None:
                        c = ws.cell(r, 5 + k, v); c.font = f_in; c.number_format = "#,##0.00"
                r += 1
            r += 1
            ws.cell(r, 1, "Year-over-year change (formula)").font = f_bold; r += 1
            yoy_row = {}
            for n, sr in enumerate(inds):
                yoy_row[n] = r
                ws.cell(r, 1, f"{sr.get('name')} - YoY %").font = Font(name=FONT, size=9)
                for k in range(4, len(qcols)):
                    a = f"{L(5 + k)}{lvl[n]}"; b = f"{L(5 + k - 4)}{lvl[n]}"
                    c = ws.cell(r, 5 + k, f'=IF(OR({a}="",{b}=""),"",{a}/{b}-1)'); c.number_format = PCT; c.font = f_norm
                r += 1
            r += 1
            ws.cell(r, 1, "Industry series: source links and notes").font = f_bold; r += 1
            head(["Series", "Source link", "Notes"])
            for sr in inds:
                us = _urls(sr.get("url")) or [""]
                for j, u in enumerate(us):
                    ws.cell(r, 1, sr.get("name") if j == 0 else "").font = Font(name=FONT, size=9)
                    if u:
                        link(r, 2, u)
                    if j == 0 and sr.get("notes"):
                        ws.cell(r, 3, sr["notes"]).font = Font(name=FONT, size=9)
                    r += 1
            # link YoY rows into the quarterly sheet
            qm = self.wb["Quarterly Model"]
            for n, rr in enumerate(getattr(self, "ind_links", [])):
                for k, q in enumerate(qcols):
                    if k < 4 or q not in hist_q:
                        continue
                    pi = hist_q.index(q)
                    c = qm.cell(rr, FC0 + pi, f"='Data Sources'!{L(5 + k)}{yoy_row[n]}"); c.font = f_link; c.number_format = PCT
            r += 2
        kp = s.get("history", {}).get("kpis") or []
        if kp:
            ws.cell(r, 1, "Company operating KPIs as reported").font = f_bold; r += 1
            head(["KPI", "Unit", "Source", ""] + [p["label"] for p in self.hist])
            for k_ in kp:
                ws.cell(r, 1, k_.get("name")).font = Font(name=FONT, size=9)
                ws.cell(r, 2, k_.get("unit")).font = Font(name=FONT, size=9)
                if k_.get("source", "").startswith("http"):
                    link(r, 3, k_["source"])
                else:
                    ws.cell(r, 3, k_.get("source")).font = Font(name=FONT, size=9)
                vals = k_.get("values") or []
                for i, v in enumerate(vals[: self.NH]):
                    v = _num(v)
                    if v is not None:
                        c = ws.cell(r, 5 + i, v); c.font = f_in
                        c.number_format = PCT if k_.get("unit") == "%" else NUM
                r += 1
            r += 2
        comps = s.get("competitors") or []
        if comps:
            ws.cell(r, 1, "Competitor data: source links").font = f_bold; r += 1
            head(["Company", "Data covered", "Source link"])
            for cp in comps:
                first = True
                for sl in cp.get("sources") or []:
                    for u in _urls(sl.get("url") if isinstance(sl, dict) else sl):
                        ws.cell(r, 1, f"{cp.get('name')} ({cp.get('ticker')})" if first else "").font = Font(name=FONT, size=9)
                        ws.cell(r, 2, sl.get("label", "") if isinstance(sl, dict) else "").font = Font(name=FONT, size=9)
                        link(r, 3, u)
                        first = False
                        r += 1
        ws.column_dimensions["A"].width = 56
        ws.column_dimensions["B"].width = 34
        ws.column_dimensions["C"].width = 60
        ws.column_dimensions["D"].width = 14


def build_workbook(spec: dict, path: str) -> dict:
    """Build the workbook at `path`; returns the validation report."""
    try:
        mb = ModelBuilder(spec)
    except SpecError as e:
        return {"ok": False, "errors": [str(e)], "warnings": [], "checks": {}}
    rep = mb.build(path)
    rep["periods"] = {"history": [mb.hist[0]["label"], mb.hist[-1]["label"]], "forecast_end": mb.periods[-1]["label"]}
    return rep
