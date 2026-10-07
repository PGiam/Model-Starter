import json
import os

import openpyxl

from app.builder.model_builder import build_workbook

HERE = os.path.dirname(__file__)


def test_chrw_spec_builds(tmp_path):
    with open(os.path.join(HERE, "..", "examples", "chrw_spec.json")) as f:
        spec = json.load(f)
    out = tmp_path / "CHRW_Model.xlsx"
    rep = build_workbook(spec, str(out))
    assert rep["ok"], rep["errors"]
    wb = openpyxl.load_workbook(out)
    for tab in ("Quarterly Model", "Annual Model", "Key Metrics", "Competitors", "Data Sources"):
        assert tab in wb.sheetnames
    # Source links live in clickable cells, never in comments.
    ds = wb["Data Sources"]
    links = [c for row in ds.iter_rows() for c in row if c.hyperlink]
    assert len(links) > 20
    assert not any(c.comment for row in ds.iter_rows() for c in row)


def test_bad_spec_reports_errors(tmp_path):
    rep = build_workbook({"company": {"ticker": "X"}}, str(tmp_path / "x.xlsx"))
    assert not rep["ok"] and rep["errors"]
